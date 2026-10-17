"""Сравнение с прошлым запуском и выгрузка в Excel."""
from __future__ import annotations

import json
from collections import defaultdict
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# ---------- сравнение с прошлым запуском ----------

CHEAPER, PRICIER, NEW, GONE, OUT, BACK = (
    "Подешевело", "Подорожало", "Новый товар", "Пропал с сайта", "Закончился", "Снова в наличии",
)


@dataclass
class Change:
    status: str
    item: dict
    old_price: float | None = None
    new_price: float | None = None

    @property
    def diff(self) -> float | None:
        if self.old_price is None or self.new_price is None:
            return None
        return round(self.new_price - self.old_price, 2)

    @property
    def diff_pct(self) -> float | None:
        if self.diff is None or not self.old_price:
            return None
        return self.diff / self.old_price


def load_snapshot(path: Path) -> dict[str, dict] | None:
    if not path.exists():
        return None
    data = json.loads(path.read_text(encoding="utf-8"))
    return {p["url"]: p for p in data["products"]}


def save_snapshot(path: Path, products: list[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    payload = {"saved_at": datetime.now().isoformat(timespec="seconds"), "products": products}
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=1), encoding="utf-8")


def compare(old: dict[str, dict] | None, new: list[dict]) -> list[Change]:
    if old is None:
        return []
    changes = []
    new_by_url = {p["url"]: p for p in new}
    for url, p in new_by_url.items():
        prev = old.get(url)
        if prev is None:
            changes.append(Change(NEW, p, None, p["price"]))
            continue
        if p["price"] < prev["price"]:
            changes.append(Change(CHEAPER, p, prev["price"], p["price"]))
        elif p["price"] > prev["price"]:
            changes.append(Change(PRICIER, p, prev["price"], p["price"]))
        if prev["in_stock"] and not p["in_stock"]:
            changes.append(Change(OUT, p, None, p["price"]))
        elif not prev["in_stock"] and p["in_stock"]:
            changes.append(Change(BACK, p, None, p["price"]))
    for url, prev in old.items():
        if url not in new_by_url:
            changes.append(Change(GONE, prev, prev["price"], None))
    order = [PRICIER, CHEAPER, NEW, GONE, OUT, BACK]
    changes.sort(key=lambda c: (order.index(c.status), -(abs(c.diff_pct or 0))))
    return changes


# ---------- оформление Excel ----------

ACCENT = "1F4E79"
HEADER_FILL = PatternFill("solid", fgColor=ACCENT)
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=16, color=ACCENT)
MUTED_FONT = Font(italic=True, color="7F7F7F")
LINK_FONT = Font(color="0563C1", underline="single")
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(bottom=THIN)
STATUS_FILL = {
    CHEAPER: PatternFill("solid", fgColor="E2EFDA"),
    PRICIER: PatternFill("solid", fgColor="FCE4D6"),
    NEW: PatternFill("solid", fgColor="DDEBF7"),
    GONE: PatternFill("solid", fgColor="EDEDED"),
    OUT: PatternFill("solid", fgColor="FFF2CC"),
    BACK: PatternFill("solid", fgColor="E2EFDA"),
}
PRICE_FMT = '#,##0 "₽"'
DIFF_FMT = '+#,##0 "₽";-#,##0 "₽";0 "₽"'

_rate = 1.0  # рублей за единицу валюты сайта, задаётся в build_report


def rub(value: float | None) -> int | None:
    """Цена в валюте сайта -> целые рубли по курсу ЦБ."""
    return None if value is None else round(value * _rate)
PCT_FMT = '+0.0%;-0.0%;0.0%'


def stars(n: int) -> str:
    return "★" * n + "☆" * (5 - n)


def write_table(ws, start_row: int, headers: list[str], widths: list[int]) -> None:
    for col, (h, w) in enumerate(zip(headers, widths), 1):
        c = ws.cell(row=start_row, column=col, value=h)
        c.fill, c.font = HEADER_FILL, HEADER_FONT
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.row_dimensions[start_row].height = 30


def sheet_products(wb: Workbook, products: list[dict], changes: list[Change]) -> None:
    ws = wb.create_sheet("Товары")
    price_change = {c.item["url"]: c for c in changes if c.status in (CHEAPER, PRICIER)}
    headers = ["№", "Название", "Категория", "Цена", "Рейтинг", "Наличие", "Изменение цены"]
    write_table(ws, 1, headers, [6, 60, 28, 12, 12, 14, 16])
    rows = sorted(products, key=lambda p: (p["category"], p["title"]))
    for i, p in enumerate(rows, 1):
        r = i + 1
        ws.cell(row=r, column=1, value=i)
        title = ws.cell(row=r, column=2, value=p["title"])
        title.hyperlink, title.font = p["url"], LINK_FONT
        ws.cell(row=r, column=3, value=p["category"])
        ws.cell(row=r, column=4, value=rub(p["price"])).number_format = PRICE_FMT
        ws.cell(row=r, column=5, value=stars(p["rating"])).font = Font(color="BF8F00")
        stock = ws.cell(row=r, column=6, value="В наличии" if p["in_stock"] else "Нет")
        if not p["in_stock"]:
            stock.font = Font(color="C00000")
        ch = price_change.get(p["url"])
        if ch:
            cell = ws.cell(row=r, column=7, value=rub(ch.diff))
            cell.number_format = DIFF_FMT
            cell.fill = STATUS_FILL[ch.status]
        for col in range(1, 8):
            ws.cell(row=r, column=col).border = BORDER
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:G{len(rows) + 1}"


def sheet_changes(wb: Workbook, changes: list[Change], first_run: bool) -> None:
    ws = wb.create_sheet("Изменения")
    if first_run:
        ws["A1"] = "Это первый запуск: сравнивать пока не с чем."
        ws["A2"] = "Запустите парсер ещё раз позже, и здесь появятся изменения цен и наличия."
        ws["A1"].font = Font(bold=True, size=12)
        ws["A2"].font = MUTED_FONT
        ws.column_dimensions["A"].width = 90
        return
    if not changes:
        ws["A1"] = "С прошлого запуска ничего не изменилось."
        ws["A1"].font = Font(bold=True, size=12)
        ws.column_dimensions["A"].width = 60
        return
    headers = ["Что произошло", "Название", "Категория", "Было", "Стало", "Разница", "Разница, %"]
    write_table(ws, 1, headers, [18, 55, 28, 12, 12, 13, 12])
    for r, c in enumerate(changes, 2):
        values = [c.status, c.item["title"], c.item["category"], rub(c.old_price), rub(c.new_price), rub(c.diff), c.diff_pct]
        for col, v in enumerate(values, 1):
            cell = ws.cell(row=r, column=col, value=v)
            cell.fill = STATUS_FILL[c.status]
            cell.border = BORDER
        ws.cell(row=r, column=2).hyperlink = c.item["url"]
        ws.cell(row=r, column=2).font = LINK_FONT
        for col in (4, 5):
            ws.cell(row=r, column=col).number_format = PRICE_FMT
        ws.cell(row=r, column=6).number_format = DIFF_FMT
        ws.cell(row=r, column=7).number_format = PCT_FMT
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:G{len(changes) + 1}"


def sheet_summary(wb: Workbook, products: list[dict], changes: list[Change], meta: dict) -> None:
    ws = wb.active
    ws.title = "Сводка"
    ws["A1"] = "Мониторинг цен"
    ws["A1"].font = TITLE_FONT
    rate_label = f"Курс ЦБ на {meta['rate_date']}" if meta["rate_date"] else "Курс задан вручную"
    ws["A2"] = (f"Сайт: {meta['site']}  ·  Обновлено: {meta['date']}  ·  "
                f"{rate_label}: 1 £ = {meta['rate']:.2f} ₽")
    ws["A2"].font = MUTED_FONT
    if meta.get("simulated"):
        ws["A3"] = "Демо-режим: часть прошлых цен изменена искусственно, чтобы показать, как выглядит отчёт об изменениях."
        ws["A3"].font = Font(italic=True, color="C00000")

    counts = defaultdict(int)
    for c in changes:
        counts[c.status] += 1
    in_stock = sum(p["in_stock"] for p in products)
    kpis = [
        ("Товаров", len(products)),
        ("В наличии", in_stock),
        ("Подорожало", counts[PRICIER]),
        ("Подешевело", counts[CHEAPER]),
        ("Новых", counts[NEW]),
        ("Пропало", counts[GONE]),
    ]
    for i, (label, value) in enumerate(kpis):
        col = 1 + i
        lab = ws.cell(row=5, column=col, value=label)
        val = ws.cell(row=6, column=col, value=value)
        lab.font = MUTED_FONT
        val.font = Font(bold=True, size=18, color=ACCENT)
        lab.alignment = val.alignment = Alignment(horizontal="center")

    by_cat: dict[str, list[float]] = defaultdict(list)
    for p in products:
        by_cat[p["category"]].append(rub(p["price"]))
    start = 9
    ws.cell(row=start - 1, column=1, value="По категориям").font = Font(bold=True, size=12)
    write_table(ws, start, ["Категория", "Товаров", "Мин. цена", "Средняя", "Макс. цена"], [30, 14, 14, 14, 14])
    for col in "ABCDEF":
        ws.column_dimensions[col].width = max(ws.column_dimensions[col].width or 0, 16)
    for r, (cat, prices) in enumerate(sorted(by_cat.items(), key=lambda kv: -len(kv[1])), start + 1):
        ws.cell(row=r, column=1, value=cat)
        ws.cell(row=r, column=2, value=len(prices))
        for col, v in ((3, min(prices)), (4, sum(prices) / len(prices)), (5, max(prices))):
            ws.cell(row=r, column=col, value=round(v)).number_format = PRICE_FMT
        for col in range(1, 6):
            ws.cell(row=r, column=col).border = BORDER


def build_report(path: Path, products: list[dict], changes: list[Change], meta: dict, first_run: bool) -> None:
    global _rate
    _rate = meta["rate"]
    wb = Workbook()
    sheet_summary(wb, products, changes, meta)
    sheet_changes(wb, changes, first_run)
    sheet_products(wb, products, changes)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
